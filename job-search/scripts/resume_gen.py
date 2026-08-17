import json, re, os, sys, collections
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from resume_content import (CONTACT, EDUCATION, SKILL_GROUPS, BBW, BBW_HEADER, GSK, GSK_HEADER)

OUT = "resumes"
os.makedirs(OUT, exist_ok=True)
NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x50, 0x50, 0x50)
FONT = "Calibri"

# ---------------- JD signal extraction -------------------------------------------
FAMILY_TERMS = {
 "sdlc": ["requirement","user stor","acceptance criteria","uat","user acceptance","test case","test plan",
          "qa","defect","sdlc","agile","scrum","backlog","sprint","traceab","business acceptance","stakeholder",
          "process flow","current state","future state","gap analysis","documentation","jira","confluence"],
 "data": ["dashboard","power bi","tableau","visuali","report","sql","python","analytics","insight","kpi",
          "metric","data analysis","etl","data set","dataset","query","excel"],
 "gov":  ["governance","data quality","stewardship","master data","metadata","lineage","catalog","compliance",
          "controls","audit","policy","standards","privacy","regulatory"],
 "tech": ["integration","api","system","architecture","technical","platform","database","backend","migration",
          "implementation","configuration","troubleshoot","solution design"],
 "cloud":["cloud","aws","azure","gcp","ci/cd","pipeline","devops","snowflake","databricks","automation"],
 "stats":["statistic","forecast","predictive","model","regression","trend","quantitative","hypothesis",
          "sas","machine learning","data science"],
 "supply":["supply chain","logistics","inventory","demand","procurement","warehouse","transportation",
           "distribution","planning","fulfillment","operations"],
 "finance":["financial","finance","revenue","cost","budget","forecast","p&l","accounting","pricing","billing",
            "reconcil","invoice","audit","epm"],
 "health":["health","patient","clinical","payer","claims","medicaid","medicare","pharmac","provider","member",
           "hipaa","care"],
}
FAMILY_LABEL = {
 "sdlc":"Requirements & Delivery","data":"Data Analytics & BI","gov":"Data Governance & Quality",
 "tech":"Systems & Technical Analysis","cloud":"Data Engineering & Cloud","stats":"Statistical Modeling",
 "supply":"Supply Chain Analytics","finance":"Financial Analysis","health":"Healthcare Data",
}
# competency phrases -> the JD signals that justify showing them (all truthful for Ahmad)
COMPETENCIES = [
 ("Requirements Elicitation & Documentation","sdlc"),("User Stories & Acceptance Criteria","sdlc"),
 ("UAT & Business Acceptance Testing","sdlc"),("Test Case Authoring & QA Partnership","sdlc"),
 ("Current/Future-State Process Analysis","sdlc"),("Requirements Traceability","sdlc"),
 ("SQL & Python Analysis","data"),("Power BI & Tableau Dashboards","data"),
 ("KPI Design & Executive Reporting","data"),("Ad-hoc & Self-Service Reporting","data"),
 ("Data Governance & Quality Controls","gov"),("Master Data, Metadata & Lineage","gov"),
 ("Data Stewardship & Auditability","gov"),("Regulatory & Compliance Reporting","gov"),
 ("Systems & Integration Impact Analysis","tech"),("ERP (SAP) Implementation Support","tech"),
 ("Solution Documentation","tech"),("ETL Pipeline Development","cloud"),
 ("Cloud Platforms & CI/CD","cloud"),("Data Modeling & Database Design","cloud"),
 ("Forecasting & Predictive Modeling","stats"),("Statistical Analysis & Regression","stats"),
 ("Capacity & Scenario Modeling","stats"),("Supply-Demand & Inventory Analytics","supply"),
 ("Logistics & Transportation Analysis","supply"),("Cost-Benefit & Margin Analysis","finance"),
 ("Financial Reporting & Reconciliation","finance"),("Healthcare & Pharmaceutical Data","health"),
 ("Stakeholder Communication","sdlc"),("Cross-Functional Collaboration","data"),
 ("Mentoring & Knowledge Transfer","sdlc"),
]

def score_jd(text):
    t = (text or "").lower()
    sc = {}
    for fam, terms in FAMILY_TERMS.items():
        sc[fam] = sum(t.count(k) for k in terms)
    return sc

def top_families(sc, n=3):
    return [f for f,_ in sorted(sc.items(), key=lambda kv:-kv[1]) if sc[f] > 0][:n]

def pick_bullets(bank, fams, want):
    ranked = []
    for txt, bf, w in bank:
        overlap = len(bf & set(fams))
        bonus = 2 if (fams and bf & {fams[0]}) else 0
        ranked.append((w*2 + overlap*3 + bonus, txt))
    ranked.sort(key=lambda x:-x[0])
    return [t for _, t in ranked[:want]]

def pick_skills(fams):
    ranked = sorted(SKILL_GROUPS, key=lambda g: -(len(g[2] & set(fams))*10 + (5 if fams and fams[0] in g[2] else 0)))
    return ranked[:6]

def pick_competencies(fams, n=12):
    out = [c for c,f in COMPETENCIES if f in fams]
    for c,f in COMPETENCIES:
        if len(out) >= n: break
        if c not in out: out.append(c)
    return out[:n]

OPENERS = {
 "sdlc": "Business analyst with 6+ years acting as the bridge between business partners and technical teams, "
         "translating business need into requirements that survive design, build, test and release.",
 "data": "Data analyst with 6+ years turning complex, high-volume enterprise datasets into decisions, owning "
         "the work end to end from extraction and modelling through dashboards and executive reporting.",
 "gov":  "Data professional with 6+ years operationalising data governance — standards, quality controls, "
         "master data and metadata — so that enterprise reporting stays trusted and auditable.",
 "tech": "Technical business analyst with 6+ years leading systems analysis across enterprise platforms, "
         "translating business need into technical requirements and integration-aware solution design.",
 "cloud":"Analytics engineer with 6+ years building ETL pipelines and automated reporting on cloud "
         "infrastructure, delivering data products that replace manual process.",
 "stats":"Analyst with 6+ years applying forecasting, regression and statistical modelling to large "
         "operational datasets, with measurable and repeatable accuracy gains.",
 "supply":"Supply chain analyst with 6+ years across inventory, demand planning, logistics and capacity, "
          "grounded in an industrial engineering degree and hands-on ERP work.",
 "finance":"Analyst with 6+ years delivering financial and operational analysis — cost-benefit, "
           "reconciliation and margin impact — to senior decision-makers.",
 "health":"Analyst with 6+ years in data-intensive roles including a global pharmaceutical program, "
          "combining health-sector domain exposure with strong SQL, BI and reporting delivery.",
}
SECOND = {
 "sdlc": "Writes user stories and acceptance criteria, authors test cases, partners with QA through defect "
         "triage, and drives UAT and business acceptance to formal sign-off.",
 "data": "Advanced SQL and Python, Power BI and Tableau, with a track record of dashboards that cut reporting "
         "time and KPI frameworks leadership actually uses.",
 "gov":  "Hands-on with data quality controls, lineage, stewardship and ERP master data, including a full "
         "SAP migration delivered at 100% data integrity.",
 "tech": "Comfortable across SAP, SQL Server, Snowflake and Informatica, performing impact and scope analysis "
         "over integrated systems before build.",
 "cloud":"Python and SQL pipelines released through CI/CD practices, plus Informatica ETL standards that make "
         "data flows consistent and auditable.",
 "stats":"Forecast accuracy improved 25% at Bath & Body Works and 20% at GSK through statistical modelling and "
         "structured performance review.",
 "supply":"Managed 1,000+ SKUs across East and West Africa and drove a $50M logistics cost reduction through "
          "cost-benefit and root cause analysis.",
 "finance":"Analysis behind $10M in annual transportation savings and a $50M logistics cost reduction, "
           "presented directly to senior leadership.",
 "health":"Managed pharmaceutical master data across 1,000+ SKUs and built the reporting that flagged supply "
          "and inventory risk in real time.",
}

def summary_for(company, jd_title, fams, sc):
    lead = fams[0] if fams else "data"
    parts = [OPENERS.get(lead, OPENERS["data"]), SECOND.get(lead, SECOND["data"])]
    for f in fams[1:3]:
        extra = SECOND.get(f)
        if extra and extra not in parts: parts.append(extra)
    parts.append("M.S. Information Systems; B.E. Industrial Engineering.")
    return " ".join(parts)

# ---------------- docx styling ---------------------------------------------------
def rule(p, color="BFBFBF", size=6):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    bt.set(qn('w:val'),'single'); bt.set(qn('w:sz'),str(size)); bt.set(qn('w:space'),'2')
    bt.set(qn('w:color'),color); b.append(bt); pPr.append(b)

def para(doc, text="", size=10.5, bold=False, color=None, space_after=2, space_before=0,
         align=None, italic=False, font=FONT):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(space_before)
    pf.line_spacing = 1.0
    if align: p.alignment = align
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        r.font.size = Pt(size); r.font.name = font
        if color is not None: r.font.color.rgb = color
    return p

def heading(doc, text):
    p = para(doc, text, size=10.5, bold=True, color=NAVY, space_before=8, space_after=3)
    for r in p.runs:
        rPr = r._r.get_or_add_rPr(); sp = OxmlElement('w:spacing'); sp.set(qn('w:val'),'30'); rPr.append(sp)
    rule(p)
    return p

def bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_after = Pt(2); pf.space_before = Pt(0); pf.line_spacing = 1.0
    pf.left_indent = Inches(0.18); pf.first_line_indent = Inches(-0.13)
    r = p.add_run(text); r.font.size = Pt(size); r.font.name = FONT
    return p

def build(job):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    for m in ("top_margin","bottom_margin"): setattr(s, m, Inches(0.5))
    for m in ("left_margin","right_margin"): setattr(s, m, Inches(0.65))

    sc   = score_jd(job["jd_text"])
    fams = top_families(sc) or ["data"]
    jd_title = (job.get("jd_title") or job["title"]).strip()

    # header
    p = para(doc, CONTACT["name"], size=20, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    for r in p.runs:
        rPr = r._r.get_or_add_rPr(); sp = OxmlElement('w:spacing'); sp.set(qn('w:val'),'60'); rPr.append(sp)
    para(doc, jd_title, size=11.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    c = f'{CONTACT["email"]}  |  {CONTACT["phone"]}  |  {CONTACT["linkedin"]}  |  {CONTACT["city"]}'
    cp = para(doc, c, size=9.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    rule(cp, color="1F3864", size=8)

    heading(doc, "PROFESSIONAL SUMMARY")
    para(doc, summary_for(job["company"], jd_title, fams, sc), size=10, space_after=3)

    heading(doc, "CORE COMPETENCIES")
    para(doc, "  ·  ".join(pick_competencies(fams)), size=9.5, space_after=3)

    heading(doc, "TECHNICAL SKILLS")
    for label, items, _ in pick_skills(fams):
        p = doc.add_paragraph()
        pf = p.paragraph_format; pf.space_after = Pt(1); pf.space_before = Pt(0); pf.line_spacing = 1.0
        r1 = p.add_run(f"{label}: "); r1.bold = True; r1.font.size = Pt(10); r1.font.name = FONT
        r2 = p.add_run(items); r2.font.size = Pt(10); r2.font.name = FONT

    heading(doc, "PROFESSIONAL EXPERIENCE")
    for (org, loc, dates), bank, want in ((BBW_HEADER, BBW, 8), (GSK_HEADER, GSK, 7)):
        role = "Senior Data Analyst / Business Analyst" if org.startswith("Bath") else \
               "Data Analyst & Reporting Engineer"
        p = doc.add_paragraph(); pf = p.paragraph_format
        pf.space_before = Pt(5); pf.space_after = Pt(0); pf.line_spacing = 1.0
        r = p.add_run(f"{role}  |  {org}"); r.bold = True; r.font.size = Pt(10.5); r.font.name = FONT
        para(doc, f"{loc}  ·  {dates}", size=9.5, color=GREY, italic=True, space_after=2)
        for b in pick_bullets(bank, fams, want): bullet(doc, b)

    heading(doc, "EDUCATION")
    for deg, school, yrs, gpa, course in EDUCATION:
        p = doc.add_paragraph(); pf = p.paragraph_format
        pf.space_before = Pt(3); pf.space_after = Pt(0); pf.line_spacing = 1.0
        r = p.add_run(f"{deg}  |  {school}"); r.bold = True; r.font.size = Pt(10); r.font.name = FONT
        para(doc, f"{yrs}  ·  {gpa}  ·  {course}", size=9.5, color=GREY, space_after=2)

    return doc, fams, sc

def slug(s, n=52):
    s = re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")
    return s[:n].rstrip("_")

jobs = json.load(open("jds.json"))
index = []
for j in jobs:
    doc, fams, sc = build(j)
    fn = f"{j['idx']:02d}_{slug(j['company'],22)}_{slug(j['title'],46)}.docx"
    doc.save(os.path.join(OUT, fn))
    index.append({"file":fn, **{k:j[k] for k in ("company","title","location","start_date","url","source","rank")},
                  "families":fams, "scores":{k:v for k,v in sorted(sc.items(), key=lambda kv:-kv[1]) if v}})
json.dump(index, open("resume_index.json","w"), indent=1)
print(f"generated {len(index)} resumes into {OUT}/")
fam_count = collections.Counter(i["families"][0] for i in index if i["families"])
print("primary JD family:", dict(fam_count))
